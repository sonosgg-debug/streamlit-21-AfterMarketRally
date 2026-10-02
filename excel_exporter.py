"""
Excel Exporter Module for AfterMarket Rally Screener
Generates formatted Excel (.xlsx) with interactive AutoFilter and numeric formats.
Complies with 00 App_AI_Template standards.
"""

import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_HEADERS


def export_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """
    스크리닝 결과를 엑셀 파일(.xlsx) 바이너리 스트림으로 생성합니다.
    - 1행 헤더에 AutoFilter(오름차순/내림차순 토글 필터) 주입
    - 블룸버그 스타일의 다크 헤더 및 깔끔한 회계/백분율 서식 적용
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "애프터마켓_급변종목_스크리닝"

    # 시트 눈금선 표시
    ws.views.sheetView[0].showGridLines = True

    # 1. 헤더 컬럼 목록 확정
    export_columns = [col for col in EXCEL_HEADERS if col in df.columns]
    if "거래소" not in export_columns and "거래소" in df.columns:
        # 4번째 위치에 거래소 삽입
        export_columns.insert(3, "거래소")

    ws.append(export_columns)

    # 2. 스타일 정의
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Malgun Gothic", size=10, bold=True, color="F8FAFC")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Malgun Gothic", size=10, color="0F172A")
    red_font = Font(name="Malgun Gothic", size=10, color="DC2626", bold=True)
    blue_font = Font(name="Malgun Gothic", size=10, color="2563EB", bold=True)
    gray_font = Font(name="Malgun Gothic", size=10, color="64748B")

    thin_border_side = Side(style="thin", color="CBD5E1")
    header_border_side = Side(style="thin", color="475569")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=header_border_side, right=header_border_side, top=header_border_side, bottom=header_border_side)

    # 헤더 행 높이 및 서식 적용
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(export_columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border

    # 3. 데이터 행 추가 및 셀 서식 적용
    for r_idx, (_, row) in enumerate(df[export_columns].iterrows(), start=2):
        ws.row_dimensions[r_idx].height = 20
        row_values = []
        for col_name in export_columns:
            val = row[col_name]
            row_values.append(val)
        ws.append(row_values)

        # 개별 셀 서식 적용
        for col_idx, col_name in enumerate(export_columns, start=1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = cell_border
            cell.font = data_font

            # 컬럼별 정렬 및 숫자 서식
            if col_name in ["종목명"]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(name="Malgun Gothic", size=10, bold=True)
            elif col_name in ["종목코드", "시장", "거래소"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "@"
            elif col_name in ["시가총액"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in ["정규장 종가(A)", "시간외 가격(B)", "정규장 거래량(C)", "시간외 거래량(D)"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            elif col_name in ["정규장 등락(%)", "시간외 등락(%)"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                try:
                    num_val = float(cell.value)
                    if num_val > 0:
                        cell.font = red_font
                        cell.number_format = '+0.00"%"'
                    elif num_val < 0:
                        cell.font = blue_font
                        cell.number_format = '-0.00"%"'
                    else:
                        cell.font = gray_font
                        cell.number_format = '0.00"%"'
                except (ValueError, TypeError):
                    pass
            elif col_name in ["시간외 거래량 비율(D/C, %)"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                try:
                    num_val = float(cell.value)
                    if num_val >= 3.0:
                        cell.font = red_font
                    cell.number_format = '0.00"%"'
                except (ValueError, TypeError):
                    pass

    # 4. 헤더 토글 기능 구현 (AutoFilter 필수 주입)
    # 사용자가 엑셀을 열었을 때 1행의 각 헤더에 오름차순/내림차순 토글 역삼각형(▼) 필터가 활성화됨
    ws.auto_filter.ref = ws.dimensions

    # 5. 열 너비 자동 최적화
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            # 한글 및 전각 문자 길이 보정 (2바이트 취급)
            length = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if length > max_len:
                max_len = length
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 6. 바이트 스트림 변환
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
